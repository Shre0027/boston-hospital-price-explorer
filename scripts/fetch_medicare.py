"""
Medicare benchmark inputs that can be downloaded directly (run on GitHub Actions).

Downloads the FY 2026 IPPS Final Rule wage index table (Table 2: wage index for each hospital,
by CMS Certification Number) and saves the wage index for our 18 hospitals to
data/medicare/wage_index.json. Outpatient (OPPS) Medicare rates for 2026 use this FY 2026 wage index.

The Medicare rate files themselves (OPPS Addendum B and the Clinical Laboratory Fee Schedule)
sit behind an AMA license agreement on cms.gov for the CPT codes they contain, so they are
downloaded by hand after accepting it, and read by scripts/medicare_rates.py.

Usage:
  python scripts/fetch_medicare.py
"""

import csv
import io
import json
import re
import time
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TABLE2_URL = "https://www.cms.gov/files/zip/fy2026-ipps-fr-tables-2-3-4a-4b.zip"


def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (boston-hospital-price-explorer)"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def rows_from(name, data):
    """Yield rows (lists of strings) from an .xlsx, .csv or .txt member."""
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                yield ["" if v is None else str(v).strip() for v in r]
    else:
        text = data.decode("utf-8-sig", "replace")
        delim = "\t" if text.count("\t") > text.count(",") else ","
        for r in csv.reader(io.StringIO(text), delimiter=delim):
            yield [v.strip() for v in r]


def main():
    hospitals = json.loads((ROOT / "config/hospitals.json").read_text())
    want = {h["ccn"]: h["id"] for h in hospitals}
    z = zipfile.ZipFile(io.BytesIO(get(TABLE2_URL)))
    names = z.namelist()
    print("::notice::Files in zip: " + "; ".join(names))
    t2 = [n for n in names if re.search(r"tables?\s*_?2\.txt$", n, re.I)] or [n for n in names if re.search(r"tables?\s*_?2\b", n, re.I) and not n.startswith("__MACOSX")]
    if not t2:
        raise SystemExit("::error::No Table 2 file found in the zip")
    found, header = {}, None
    for name in t2:
        for r in rows_from(name, z.read(name)):
            joined = " | ".join(r).lower()
            if header is None and sum(1 for v in r if v) >= 5 and ("provider" in joined or "ccn" in joined) and "wage index" in joined:
                header = r
                print(f"::notice::{name} header: " + " | ".join(header))
                continue
            if header and r and r[0].zfill(6) in want:
                found[want[r[0].zfill(6)]] = dict(zip(header, r))
    print(f"::notice::Found {len(found)} of {len(want)} hospitals in Table 2")
    for hid, rec in found.items():
        print(f"::notice::{hid}: " + "; ".join(f"{k}={v}" for k, v in rec.items() if v)[:300])
    out = ROOT / "data/medicare"
    out.mkdir(parents=True, exist_ok=True)
    (out / "wage_index.json").write_text(json.dumps({"source": TABLE2_URL, "fetched_on": time.strftime("%Y-%m-%d", time.gmtime()),
                                                     "header": header, "hospitals": found}, indent=1))


if __name__ == "__main__":
    main()
