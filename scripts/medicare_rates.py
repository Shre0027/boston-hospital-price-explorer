"""
Turn the official Medicare rate files into a small benchmark table for our procedures.

Inputs (downloaded by hand from cms.gov after accepting the AMA CPT license; NOT committed,
because the license restricts redistributing the files and their code descriptions):
  - OPPS Addendum B zip (outpatient payment rates by HCPCS code), e.g. October-2026-OPPS-Addendum-B.zip
  - Clinical Laboratory Fee Schedule zip, e.g. 26CLABQ4.zip
  - data/medicare/wage_index.json (from scripts/fetch_medicare.py)

Output: data/medicare/rates.json, containing only code numbers and Medicare's dollar rates (no descriptions).

Method (Medicare's own payment formulas):
  - Outpatient services: national unadjusted OPPS payment rate x (0.60 x wage index + 0.40).
    0.60 is the OPPS labor-related share. All 18 hospitals have the same FY 2026 wage index.
  - Lab tests: the Clinical Laboratory Fee Schedule national limitation amount (not wage adjusted).
  - Not benchmarked: services Medicare doesn't pay as a separate outpatient rate (status indicators
    such as N, Q1-Q4 packaged, or A = paid under another fee schedule), the office visit (99213; the
    hospital-side equivalent is a different code, G0463), and the inpatient gallbladder stay (MS-DRG 419,
    paid under the inpatient system with hospital-specific add-ons).

Usage:
  python scripts/medicare_rates.py --opps PATH_TO_ADDENDUM_B.zip --clfs PATH_TO_CLFS.zip
"""

import argparse
import csv
import io
import json
import re
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LABOR_SHARE = 0.60
OPPS_CODES = ["73721", "72148", "70551", "74177", "71046", "76700", "77067", "45378", "66984", "29881", "93005", "99283", "99284"]
CLFS_CODES = ["85025", "80053", "80061"]


def table_rows(path):
    """All rows from every .xlsx/.csv/.txt inside a zip (lists of strings)."""
    z = zipfile.ZipFile(path)
    for name in z.namelist():
        if name.startswith("__MACOSX") or name.endswith("/"):
            continue
        data = z.read(name)
        if name.lower().endswith(".zip"):           # CMS sometimes nests zips
            nested = Path("/tmp") / Path(name).name
            nested.write_bytes(data)
            yield from table_rows(nested)
        elif name.lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for ws in wb.worksheets:
                for r in ws.iter_rows(values_only=True):
                    yield name, ["" if v is None else str(v).strip() for v in r]
        elif name.lower().endswith((".csv", ".txt")):
            text = data.decode("utf-8-sig", "replace")
            delim = "\t" if text.count("\t") > text.count(",") else ","
            for r in csv.reader(io.StringIO(text), delimiter=delim):
                yield name, [v.strip() for v in r]


def money(v):
    v = re.sub(r"[$,\s]", "", v or "")
    try:
        return float(v)
    except ValueError:
        return None


def find_header(rows, must):
    """Return (index map, remaining rows) once a row containing all `must` patterns is seen."""
    for name, r in rows:
        low = [c.lower() for c in r]
        idx = {}
        for key, pat in must.items():
            hit = [i for i, c in enumerate(low) if re.search(pat, c)]
            if not hit:
                break
            idx[key] = hit[0]
        else:
            return name, idx
    raise SystemExit(f"::error::Could not find a header with {list(must)}")


def read_opps(path):
    rows = table_rows(path)
    name, idx = find_header(rows, {"code": r"^hcpcs", "si": r"^si$|status ind", "rate": r"payment\s*rate"})
    print(f"Addendum B: {name}, columns {idx}")
    out = {}
    for _, r in rows:
        if len(r) <= max(idx.values()):
            continue
        code = r[idx["code"]].strip().upper()
        if code in OPPS_CODES:
            out[code] = {"status_indicator": r[idx["si"]].strip(), "national_rate": money(r[idx["rate"]])}
    return out


def read_clfs(path):
    rows = table_rows(path)
    name, idx = find_header(rows, {"code": r"^hcpcs|^code", "rate": r"^rate$|national limit|nla"})
    print(f"CLFS: {name}, columns {idx}")
    out = {}
    for _, r in rows:
        if len(r) <= max(idx.values()):
            continue
        code = r[idx["code"]].strip().upper()
        if code in CLFS_CODES and code not in out:      # first row = no modifier
            out[code] = {"national_limit": money(r[idx["rate"]])}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--opps", required=True)
    ap.add_argument("--clfs", required=True)
    a = ap.parse_args()
    wi_data = json.loads((ROOT / "data/medicare/wage_index.json").read_text())
    col = next(k for k in wi_data["header"] if "2026 wage index with cap" in k.lower())
    wis = {hid: float(rec[col]) for hid, rec in wi_data["hospitals"].items()}
    if len(set(wis.values())) != 1:
        raise SystemExit(f"::error::Hospitals have different wage indexes {wis}; per-hospital rates needed")
    wi = next(iter(wis.values()))
    factor = LABOR_SHARE * wi + (1 - LABOR_SHARE)
    opps, clfs = read_opps(a.opps), read_clfs(a.clfs)
    rates = {}
    for code in OPPS_CODES:
        o = opps.get(code)
        if not o:
            rates[code] = {"system": "OPPS", "benchmark": None, "why_not": "code not in Addendum B"}
        elif o["national_rate"] is None or o["status_indicator"].upper() in {"N", "Q1", "Q2", "Q3", "Q4", "A", "B", "M"}:
            rates[code] = {"system": "OPPS", "status_indicator": o["status_indicator"], "benchmark": None,
                           "why_not": f"no separate outpatient payment rate (status indicator {o['status_indicator']})"}
        else:
            rates[code] = {"system": "OPPS", "status_indicator": o["status_indicator"], "national_rate": o["national_rate"],
                           "benchmark": round(o["national_rate"] * factor, 2)}
    for code in CLFS_CODES:
        c = clfs.get(code)
        rates[code] = {"system": "CLFS", "benchmark": c["national_limit"] if c else None, **({} if c else {"why_not": "code not in CLFS file"})}
    for code, r in rates.items():
        print(f"{code}: {r}")
    out = {"made_on": time.strftime("%Y-%m-%d", time.gmtime()), "opps_file": Path(a.opps).name, "clfs_file": Path(a.clfs).name,
           "wage_index": wi, "labor_share": LABOR_SHARE, "wage_adjustment_factor": round(factor, 5), "rates": rates,
           "not_benchmarked": {"99213": "hospital outpatient clinic visits are paid under a different code (G0463)",
                               "419": "inpatient stay, paid under the inpatient system with hospital-specific add-ons"}}
    (ROOT / "data/medicare/rates.json").write_text(json.dumps(out, indent=1))
    print("Wrote data/medicare/rates.json")


if __name__ == "__main__":
    main()
